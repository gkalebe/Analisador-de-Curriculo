import subprocess

import docx
import pytest

from app.adapters.ai_service.curriculo_schema import DadosCurriculoEstruturado, ExperienciaCurriculo
from app.adapters.curriculo_exporter.curriculo_template_exporter import (
    ConversaoPdfError,
    MontadorDocumentoCurriculo,
    TemplateCurriculoNaoEncontradoError,
)


def test_gerar_docx_template_moderno_preenche_campos(tmp_path):
    dados = DadosCurriculoEstruturado(
        nome_completo="Ana Beatriz Souza",
        titulo_profissional="Desenvolvedora Backend Pleno",
        experiencias=[ExperienciaCurriculo(cargo="Dev", empresa="XPTO", descricao_bullets=["Fez X"])],
    )
    montador = MontadorDocumentoCurriculo()

    caminho = montador.gerar_docx(dados, "moderno", str(tmp_path / "saida.docx"))

    documento = docx.Document(caminho)
    texto_completo = "\n".join(p.text for p in documento.paragraphs)
    assert "Ana Beatriz Souza" in texto_completo
    assert "Desenvolvedora Backend Pleno" in texto_completo
    assert "XPTO" in texto_completo
    assert "Fez X" in texto_completo


def test_gerar_docx_template_classico_tambem_funciona(tmp_path):
    dados = DadosCurriculoEstruturado(nome_completo="Carlos Lima")
    montador = MontadorDocumentoCurriculo()

    caminho = montador.gerar_docx(dados, "classico", str(tmp_path / "saida.docx"))

    documento = docx.Document(caminho)
    assert "Carlos Lima" in documento.paragraphs[0].text


def test_gerar_docx_com_secoes_vazias_nao_quebra_e_oculta_secao(tmp_path):
    dados = DadosCurriculoEstruturado(nome_completo="Sem Experiência")
    montador = MontadorDocumentoCurriculo()

    caminho = montador.gerar_docx(dados, "moderno", str(tmp_path / "saida.docx"))

    documento = docx.Document(caminho)
    texto_completo = "\n".join(p.text for p in documento.paragraphs)
    assert "EXPERIÊNCIA PROFISSIONAL" not in texto_completo


def test_gerar_docx_com_dict_cru_faltando_campos_usa_fallback_get(tmp_path):
    montador = MontadorDocumentoCurriculo()

    # simula uma fonte externa que pulou a Etapa 2 e manda um dict incompleto
    caminho = montador.gerar_docx({"nome_completo": "Fulano"}, "moderno", str(tmp_path / "saida.docx"))

    documento = docx.Document(caminho)
    assert "Fulano" in documento.paragraphs[0].text


def test_gerar_docx_com_template_inexistente_lanca_erro(tmp_path):
    montador = MontadorDocumentoCurriculo()

    with pytest.raises(TemplateCurriculoNaoEncontradoError):
        montador.gerar_docx(DadosCurriculoEstruturado(), "inexistente", str(tmp_path / "saida.docx"))


def test_converter_para_pdf_sem_libreoffice_instalado_lanca_erro(tmp_path):
    (tmp_path / "entrada.docx").write_bytes(b"conteudo qualquer")
    montador = MontadorDocumentoCurriculo(caminho_libreoffice="binario-que-nao-existe-no-path")

    with pytest.raises(ConversaoPdfError):
        montador.converter_para_pdf(str(tmp_path / "entrada.docx"))


def test_converter_para_pdf_com_falha_no_processo_lanca_erro(tmp_path, monkeypatch):
    (tmp_path / "entrada.docx").write_bytes(b"conteudo qualquer")
    montador = MontadorDocumentoCurriculo(caminho_libreoffice="soffice-falso")

    def run_falso(*args, **kwargs):
        return subprocess.CompletedProcess(args, returncode=1, stdout="", stderr="erro simulado do LibreOffice")

    monkeypatch.setattr(subprocess, "run", run_falso)

    with pytest.raises(ConversaoPdfError):
        montador.converter_para_pdf(str(tmp_path / "entrada.docx"))


def test_converter_para_pdf_com_sucesso_retorna_caminho(tmp_path, monkeypatch):
    (tmp_path / "entrada.docx").write_bytes(b"conteudo qualquer")
    montador = MontadorDocumentoCurriculo(caminho_libreoffice="soffice-falso")

    def run_falso(comando, **kwargs):
        (tmp_path / "entrada.pdf").write_bytes(b"%PDF-1.4 conteudo falso")
        return subprocess.CompletedProcess(comando, returncode=0, stdout="", stderr="")

    monkeypatch.setattr(subprocess, "run", run_falso)

    caminho_pdf = montador.converter_para_pdf(str(tmp_path / "entrada.docx"), str(tmp_path))

    assert caminho_pdf.endswith("entrada.pdf")
