import io

import docx
import pytest

from app.adapters.curriculo_exporter.curriculo_exporter import CurriculoExporter, TemplateNaoSuportadoError

DADOS_ESTRUTURADOS = {
    "nome": "Ana Silva",
    "email": "ana@email.com",
    "telefone": "(61) 99999-0000",
    "resumo": "Resumo profissional de duas frases sobre a candidata.",
    "formacao": "Ciência da Computação — UCB (cursando)\nTécnico em Informática — IFB (concluído)",
    "experiencia_profissional": "Analista — Empresa X (2024-atual)\nEstagiária — Empresa Y (2023-2024)",
    "habilidades": "Python, FastAPI, React, SQL",
}

DADOS_FALLBACK = {
    "nome": "",
    "email": "",
    "telefone": "",
    "resumo": "",
    "formacao": "",
    "experiencia_profissional": "",
    "habilidades": "",
    "texto_bruto": "Currículo em texto corrido, com “aspas curvas”, travessão — e reticências…",
}


@pytest.fixture
def exporter() -> CurriculoExporter:
    return CurriculoExporter()


@pytest.mark.parametrize("id_template", sorted(CurriculoExporter.TEMPLATES_SUPORTADOS))
@pytest.mark.parametrize("dados", [DADOS_ESTRUTURADOS, DADOS_FALLBACK], ids=["estruturado", "fallback"])
def test_gerar_pdf_produz_pdf_valido_para_todo_template(exporter, dados, id_template):
    conteudo = exporter.gerar_pdf(dados, id_template)

    assert conteudo.startswith(b"%PDF")
    assert len(conteudo) > 500


@pytest.mark.parametrize("id_template", sorted(CurriculoExporter.TEMPLATES_SUPORTADOS))
@pytest.mark.parametrize("dados", [DADOS_ESTRUTURADOS, DADOS_FALLBACK], ids=["estruturado", "fallback"])
def test_gerar_docx_produz_docx_valido_para_todo_template(exporter, dados, id_template):
    conteudo = exporter.gerar_docx(dados, id_template)

    assert conteudo[:2] == b"PK"
    documento = docx.Document(io.BytesIO(conteudo))
    texto_completo = "\n".join(paragrafo.text for paragrafo in documento.paragraphs)
    if dados.get("nome"):
        assert dados["nome"] in texto_completo


@pytest.mark.parametrize("id_template", sorted(CurriculoExporter.TEMPLATES_SUPORTADOS))
def test_gerar_docx_usa_marcadores_reais_para_experiencia_e_formacao(exporter, id_template):
    conteudo = exporter.gerar_docx(DADOS_ESTRUTURADOS, id_template)

    documento = docx.Document(io.BytesIO(conteudo))
    estilos_usados = {paragrafo.style.name for paragrafo in documento.paragraphs}
    assert "List Bullet" in estilos_usados


def test_gerar_pdf_com_template_inexistente_lanca_erro(exporter):
    with pytest.raises(TemplateNaoSuportadoError):
        exporter.gerar_pdf(DADOS_ESTRUTURADOS, "inexistente")


def test_gerar_docx_com_template_inexistente_lanca_erro(exporter):
    with pytest.raises(TemplateNaoSuportadoError):
        exporter.gerar_docx(DADOS_ESTRUTURADOS, "inexistente")


def test_gerar_pdf_nao_quebra_com_lista_longa_de_habilidades(exporter):
    dados = dict(DADOS_ESTRUTURADOS)
    dados["habilidades"] = ", ".join(f"HabilidadeComNomeBemLongoNumero{i}" for i in range(30))

    for id_template in CurriculoExporter.TEMPLATES_SUPORTADOS:
        conteudo = exporter.gerar_pdf(dados, id_template)
        assert conteudo.startswith(b"%PDF")


def test_dividir_itens_ignora_marcadores_e_linhas_vazias(exporter):
    texto = "- Item um\n• Item dois\n\n  Item três  "

    assert exporter._dividir_itens(texto) == ["Item um", "Item dois", "Item três"]


def test_dividir_habilidades_separa_por_virgula_e_ignora_vazios(exporter):
    texto = "Python, , FastAPI ,SQL"

    assert exporter._dividir_habilidades(texto) == ["Python", "FastAPI", "SQL"]
